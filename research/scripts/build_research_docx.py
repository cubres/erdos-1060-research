#!/usr/bin/env python3
"""Build the Word edition of the research note."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/cubres/Documents/ChatGPT/Research")
SOURCE = ROOT / "report-source.md"
OUTPUT_DIR = ROOT / "output"
OUTPUT = OUTPUT_DIR / "Erdos_Problem_1060_Research_Note.docx"
BUILD_DIR = ROOT / "tmp" / "docx_build"
MATH_DIR = BUILD_DIR / "math"
MANIFEST = BUILD_DIR / "math_manifest.json"
TEX_BIN = Path("/Users/cubres/Library/TeXLive/2026/bin/universal-darwin")
LATEX = TEX_BIN / "latex"
DVIPNG = TEX_BIN / "dvipng"

NAVY = "17324D"
PALE_BLUE = "EEF4F8"
LIGHT_GRAY = "D9D9D9"
MID_GRAY = RGBColor(90, 100, 110)
LINK_BLUE = "1F5A91"


def set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn("w:" + edge))
        if tag is None:
            tag = OxmlElement("w:" + edge)
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), LIGHT_GRAY)


def set_picture_alt(run, description: str) -> None:
    drawing = run._element.find(qn("w:drawing"))
    if drawing is None:
        return
    doc_pr = drawing.find(".//wp:docPr", namespaces=drawing.nsmap)
    if doc_pr is not None:
        doc_pr.set("descr", description)
        doc_pr.set("title", "Mathematical equation")


def add_hyperlink(paragraph, label: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = label
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def math_id(tex: str, display: bool) -> str:
    key = ("D:" if display else "I:") + tex
    return hashlib.sha1(key.encode("utf-8")).hexdigest()


def collect_math(markdown: str) -> list[dict]:
    markdown = "\n".join(
        re.sub(r"^\s*>\s?", "", line) for line in markdown.splitlines()
    )
    found: dict[str, dict] = {}
    for match in re.finditer(r"\$\$(.*?)\$\$", markdown, flags=re.S):
        tex = match.group(1).strip()
        ident = math_id(tex, True)
        found[ident] = {"id": ident, "tex": tex, "display": True}
    without_display = re.sub(r"\$\$(.*?)\$\$", "", markdown, flags=re.S)
    pattern = r"(?<!\$)\$(?!\$)(.*?)(?<!\$)\$(?!\$)"
    for match in re.finditer(pattern, without_display):
        tex = match.group(1).strip()
        ident = math_id(tex, False)
        found[ident] = {"id": ident, "tex": tex, "display": False}
    return list(found.values())


def render_math(markdown: str) -> None:
    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    MATH_DIR.mkdir(parents=True, exist_ok=True)
    items = collect_math(markdown)
    MANIFEST.write_text(json.dumps(items, ensure_ascii=False, indent=2))
    latex_lines = [
        r"\documentclass{article}",
        r"\usepackage[active,tightpage]{preview}",
        r"\usepackage{amsmath,amssymb}",
        r"\PreviewBorder=2pt",
        r"\pagestyle{empty}",
        r"\begin{document}",
    ]
    for item in items:
        style = r"\displaystyle " if item["display"] else ""
        latex_lines.extend(
            [
                r"\begin{preview}",
                r"\(" + style + item["tex"] + r"\)",
                r"\end{preview}",
            ]
        )
    latex_lines.append(r"\end{document}")
    tex_path = BUILD_DIR / "math.tex"
    tex_path.write_text("\n".join(latex_lines))
    subprocess.run(
        [
            str(LATEX),
            "-interaction=nonstopmode",
            "-halt-on-error",
            "-output-directory",
            str(BUILD_DIR),
            str(tex_path),
        ],
        check=True,
        cwd=BUILD_DIR,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )
    page_pattern = BUILD_DIR / "equation-%04d.png"
    subprocess.run(
        [
            str(DVIPNG),
            "-D",
            "420",
            "-bg",
            "Transparent",
            "-T",
            "tight",
            "-o",
            str(page_pattern),
            str(BUILD_DIR / "math.dvi"),
        ],
        check=True,
        cwd=BUILD_DIR,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )
    rendered = sorted(BUILD_DIR.glob("equation-*.png"))
    if len(rendered) != len(items):
        raise RuntimeError(
            "Equation page mismatch: rendered "
            + str(len(rendered))
            + " expected "
            + str(len(items))
        )
    for item, source_path in zip(items, rendered):
        shutil.copyfile(source_path, MATH_DIR / (item["id"] + ".png"))


INLINE_TOKEN = re.compile(
    r"(\*\*.*?\*\*|(?<!\*)\*[^*\n]+?\*(?!\*)|\[[^\]]+\]\([^)]+\)|\x60[^\x60]+\x60|"
    r"(?<!\$)\$(?!\$).*?(?<!\$)\$(?!\$))"
)


def add_inline(paragraph, text: str, default_size: float = 10.7) -> None:
    position = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, "Cambria", default_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            set_run_font(run, "Cambria", default_size)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            set_run_font(run, "Cambria", default_size)
        elif token.startswith("["):
            link_match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link_match:
                add_hyperlink(paragraph, link_match.group(1), link_match.group(2))
        elif token.startswith(chr(96)):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Aptos Mono", max(8.8, default_size - 0.7))
            run.font.color.rgb = RGBColor(45, 55, 65)
        elif token.startswith("$"):
            tex = token[1:-1].strip()
            image_path = MATH_DIR / (math_id(tex, False) + ".png")
            run = paragraph.add_run()
            run.add_picture(str(image_path), height=Pt(default_size + 2.4))
            set_picture_alt(run, tex)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, "Cambria", default_size)


def add_display_equation(doc: Document, tex: str, quote: bool = False) -> None:
    image_path = MATH_DIR / (math_id(tex, True) + ".png")
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    aspect = width_px / max(height_px, 1)
    line_count = max(1, tex.count(r"\\") + 1 if "aligned" in tex else 1)
    desired_height = 0.26 + 0.23 * (line_count - 1)
    desired_width = desired_height * aspect
    max_width = 5.9 if quote else 6.55
    if desired_width > max_width:
        desired_width = max_width
        desired_height = desired_width / aspect
    if desired_width < 0.8:
        desired_width = 0.8
        desired_height = desired_width / aspect
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if quote:
        paragraph.paragraph_format.left_indent = Inches(0.3)
        paragraph.paragraph_format.right_indent = Inches(0.2)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run()
    run.add_picture(
        str(image_path),
        width=Inches(desired_width),
        height=Inches(desired_height),
    )
    set_picture_alt(run, tex)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Cambria"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
    normal.font.size = Pt(10.7)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Cambria"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 15.0, 16, 7),
        ("Heading 2", 12.2, 11, 5),
        ("Heading 3", 11.0, 9, 4),
    ):
        style = styles[name]
        style.font.name = "Cambria"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Cambria"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
        style.font.size = Pt(10.7)
        style.paragraph_format.space_after = Pt(3)


def add_footer(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Erdős Problem 1060 Research Note   ")
    set_run_font(run, "Cambria", 8.2)
    run.font.color.rgb = MID_GRAY
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_separate = OxmlElement("w:fldChar")
    field_separate.set(qn("w:fldCharType"), "separate")
    page_text = OxmlElement("w:t")
    page_text.text = "1"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run2 = paragraph.add_run()
    run2._r.extend([field_begin, instruction, field_separate, page_text, field_end])
    set_run_font(run2, "Cambria", 8.2)
    run2.font.color.rgb = MID_GRAY


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = len(rows[0])
    table = doc.add_table(rows=1, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths_by_count = {
        2: [2.2, 4.25],
        3: [1.05, 1.55, 3.85],
        4: [0.95, 2.28, 0.95, 2.27],
    }
    widths = widths_by_count.get(column_count, [6.45 / column_count] * column_count)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(rows[0]):
        cell = header.cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        set_cell_shading(cell, NAVY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if len(value) < 24 else WD_ALIGN_PARAGRAPH.LEFT
        )
        paragraph.paragraph_format.space_after = Pt(0)
        add_inline(paragraph, value, default_size=9.2)
        for run in paragraph.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, source_row in enumerate(rows[1:], start=1):
        row = table.add_row()
        for index, value in enumerate(source_row):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if index == 0 and len(value) < 18
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            add_inline(paragraph, value, default_size=9.1)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def is_special(line: str) -> bool:
    stripped = line.strip()
    return (
        not stripped
        or stripped.startswith("#")
        or stripped.startswith("$$")
        or stripped.startswith(chr(96) * 3)
        or stripped.startswith(">")
        or stripped.startswith("|")
        or bool(re.match(r"^[-*]\s+", stripped))
        or bool(re.match(r"^\d+\.\s+", stripped))
    )


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        parts = [part.strip() for part in lines[index].strip().strip("|").split("|")]
        if index == start + 1 and all(re.fullmatch(r":?-{3,}:?", part) for part in parts):
            index += 1
            continue
        rows.append(parts)
        index += 1
    return rows, index


def add_block_lines(doc: Document, lines: list[str], quote: bool = False) -> None:
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped == "$$":
            equation_lines = []
            index += 1
            while index < len(lines) and lines[index].strip() != "$$":
                equation_lines.append(lines[index])
                index += 1
            add_display_equation(doc, "\n".join(equation_lines).strip(), quote=quote)
            index += 1
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines) and lines[index].strip() and lines[index].strip() != "$$":
            paragraph_lines.append(lines[index].strip())
            index += 1
        paragraph = doc.add_paragraph()
        if quote:
            paragraph.paragraph_format.left_indent = Inches(0.35)
            paragraph.paragraph_format.right_indent = Inches(0.2)
        paragraph.paragraph_format.space_after = Pt(6)
        add_inline(paragraph, " ".join(paragraph_lines))


def build() -> None:
    markdown = SOURCE.read_text()
    render_math(markdown)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.67)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.32)
    add_footer(section)

    lines = markdown.splitlines()
    index = 0
    first_h2 = True
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            add_inline(paragraph, stripped[2:].strip(), default_size=24)
            for run in paragraph.runs:
                run.font.bold = True
            index += 1
            continue
        if stripped.startswith("## "):
            title = stripped[3:].strip()
            if first_h2:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(8)
                run = paragraph.add_run(title)
                set_run_font(run, "Cambria", 12.2)
                run.italic = True
                run.font.color.rgb = MID_GRAY
                first_h2 = False
            else:
                paragraph = doc.add_paragraph(title, style="Heading 1")
            index += 1
            continue
        if stripped.startswith("### "):
            doc.add_paragraph(stripped[4:].strip(), style="Heading 2")
            index += 1
            continue
        if stripped == "$$":
            equation_lines = []
            index += 1
            while index < len(lines) and lines[index].strip() != "$$":
                equation_lines.append(lines[index])
                index += 1
            add_display_equation(doc, "\n".join(equation_lines).strip())
            index += 1
            continue
        if stripped.startswith(chr(96) * 3):
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith(chr(96) * 3):
                code_lines.append(lines[index].rstrip())
                index += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.right_indent = Inches(0.2)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(7)
            for line_number, code_line in enumerate(code_lines):
                if line_number:
                    paragraph.add_run().add_break()
                run = paragraph.add_run(code_line)
                set_run_font(run, "Aptos Mono", 8.8)
                run.font.color.rgb = RGBColor(45, 55, 65)
            index += 1
            continue
        if stripped.startswith(">"):
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_line = lines[index].strip()[1:]
                if quote_line.startswith(" "):
                    quote_line = quote_line[1:]
                quote_lines.append(quote_line)
                index += 1
            add_block_lines(doc, quote_lines, quote=True)
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue
        bullet_match = re.match(r"^[-*]\s+(.*)$", stripped)
        number_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if bullet_match or number_match:
            if bullet_match:
                paragraph = doc.add_paragraph(style="List Bullet")
                add_inline(paragraph, bullet_match.group(1))
            else:
                # Preserve the source numbering exactly.  Word's built-in list
                # style otherwise continues numbering across unrelated lists.
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.25)
                paragraph.paragraph_format.first_line_indent = Inches(-0.20)
                paragraph.paragraph_format.space_after = Pt(3)
                prefix = paragraph.add_run(number_match.group(1) + ". ")
                set_run_font(prefix, "Cambria", 10.7)
                add_inline(paragraph, number_match.group(2))
            index += 1
            continue
        paragraph_lines = [stripped.rstrip()]
        index += 1
        while index < len(lines) and not is_special(lines[index]):
            paragraph_lines.append(lines[index].strip())
            index += 1
        text = " ".join(item[:-2] if item.endswith("  ") else item for item in paragraph_lines)
        paragraph = doc.add_paragraph()
        if text == "Prepared for mathematical review 5 September 2026":
            paragraph.paragraph_format.space_after = Pt(12)
            run = paragraph.add_run("Prepared for mathematical review\n5 September 2026")
            set_run_font(run, "Cambria", 9.4)
            run.font.color.rgb = MID_GRAY
        else:
            add_inline(paragraph, text)

    properties = doc.core_properties
    properties.title = "Erdős Problem 1060 Research Note"
    properties.subject = "Uniform bounds and exact collisions for k sigma(k)"
    properties.keywords = "number theory, divisor sum, sigma function, Erdős Problem 1060"
    properties.author = "Research note prepared for mathematical review"
    properties.comments = "Generated from report-source.md with exact certificates."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
